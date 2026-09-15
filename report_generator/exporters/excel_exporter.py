import pandas as pd
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from .format_utils import format_for_display


def export_excel(totals_df: pd.DataFrame, trend_df: pd.DataFrame, metadata: dict, output_path: str):
    sheet_name = metadata["period_label"][:31]  # Excel sheet name limit
    header_rows = 4  # metadata lines written above the table

    display_totals = format_for_display(totals_df.drop(columns=["is_subtotal"], errors="ignore"))

    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        display_totals.to_excel(writer, sheet_name=sheet_name, index=False, startrow=header_rows)
        worksheet = writer.sheets[sheet_name]

        worksheet.cell(row=1, column=1, value=f"{metadata['title']} — {metadata['period_label']}").font = Font(
            bold=True, size=13
        )
        worksheet.cell(row=2, column=1, value=f"Period: {metadata['period_start']} to {metadata['period_stop']}")
        worksheet.cell(row=3, column=1, value=f"Metric: {metadata['metric_name']}    Realm: {metadata['realm']}")
        generated_line = f"Generated: {metadata['generated_at']}"
        if metadata.get("resolution_ms"):
            generated_line += f"    Data resolution: ~{metadata['resolution_ms'] // 60000} min"
        worksheet.cell(row=4, column=1, value=generated_line)

        for col_idx, col_name in enumerate(display_totals.columns, start=1):
            width = max(len(str(col_name)) + 2, display_totals[col_name].astype(str).str.len().max() + 2)
            worksheet.column_dimensions[get_column_letter(col_idx)].width = width

        if "is_subtotal" in totals_df.columns:
            n_cols = len(display_totals.columns)
            for i, is_sub in enumerate(totals_df["is_subtotal"]):
                if is_sub:
                    row_idx = header_rows + 2 + i  # +1 for header row itself, +1 for 1-indexing
                    for col_idx in range(1, n_cols + 1):
                        worksheet.cell(row=row_idx, column=col_idx).font = Font(bold=True)

        if trend_df is not None and not trend_df.empty:
            trend_display = trend_df.copy()
            trend_display["period_start"] = trend_display["period_start"].dt.strftime("%Y-%m-%d")
            trend_display = format_for_display(trend_display)
            trend_display.to_excel(writer, sheet_name="Trend", index=False)
            trend_sheet = writer.sheets["Trend"]
            for col_idx, col_name in enumerate(trend_display.columns, start=1):
                width = max(len(str(col_name)) + 2, trend_display[col_name].astype(str).str.len().max() + 2)
                trend_sheet.column_dimensions[get_column_letter(col_idx)].width = width
